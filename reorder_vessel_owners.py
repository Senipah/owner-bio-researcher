from __future__ import annotations

import argparse
import sys
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, urlparse

from src.auth import authenticated_context
from src.io_utils import atomic_write_json, utc_now
from src.vessel_owner_order import (
    build_vessel_owner_order_plan,
    fetch_vessel_owner_order,
    submit_vessel_owner_order,
    vessel_edit_url,
)


def _stamp() -> str:
    return datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Dry-run and optionally repair vessel UBO relationship order "
            "using the site's oldest-first date ordering."
        )
    )
    parser.add_argument(
        "--input",
        type=Path,
        required=True,
        help="Candidate .xlsx file containing priority and vessel_id columns.",
    )
    parser.add_argument(
        "--apply",
        action="store_true",
        help="Save and verify planned order changes. Default is dry-run.",
    )
    parser.add_argument(
        "--priority",
        action="append",
        type=int,
        dest="priorities",
        help="Only process this priority; may be supplied more than once.",
    )
    parser.add_argument(
        "--vessel-id",
        action="append",
        type=int,
        dest="vessel_ids",
        help="Only process this vessel ID; may be supplied more than once.",
    )
    parser.add_argument(
        "--limit",
        type=int,
        help="Maximum number of selected workbook rows to process.",
    )
    parser.add_argument("--headless", action="store_true")
    parser.add_argument(
        "--audit-dir", type=Path, default=Path("output/audits")
    )
    return parser


def _positive_integer(value: object, *, field: str, row_number: int) -> int:
    if isinstance(value, bool):
        raise ValueError(f"Row {row_number}: {field} must be a positive integer")
    if isinstance(value, int):
        result = value
    elif isinstance(value, float) and value.is_integer():
        result = int(value)
    elif isinstance(value, str) and value.strip().isdigit():
        result = int(value.strip())
    else:
        raise ValueError(f"Row {row_number}: {field} must be a positive integer")
    if result <= 0:
        raise ValueError(f"Row {row_number}: {field} must be a positive integer")
    return result


def _link_vessel_id(value: object) -> int | None:
    link = str(value or "").strip()
    if not link:
        return None
    parsed = urlparse(link)
    if parsed.netloc.casefold() != "agent.superyachtnetwork.com":
        raise ValueError(f"Unexpected vessel link host: {parsed.netloc!r}")
    if parsed.path != "/vessel/edit/specification_new/edit.htm":
        raise ValueError(f"Unexpected vessel link path: {parsed.path!r}")
    values = parse_qs(parsed.query).get("id", [])
    if len(values) != 1 or not values[0].isdigit():
        raise ValueError(f"Vessel link has no valid ID: {link}")
    return int(values[0])


def load_vessel_candidates(path: str | Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl is required; recreate the venv from requirements.txt"
        ) from exc

    source = Path(path)
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        try:
            raw_headers = next(rows)
        except StopIteration as exc:
            raise ValueError("Candidate workbook is empty") from exc
        headers = [str(value or "").strip().casefold() for value in raw_headers]
        required = {"priority", "vessel_id", "vessel_name"}
        missing = sorted(required - set(headers))
        if missing:
            raise ValueError(
                "Candidate workbook is missing columns: " + ", ".join(missing)
            )
        column = {name: headers.index(name) for name in set(headers) if name}

        candidates: list[dict[str, Any]] = []
        seen_ids: set[int] = set()
        for row_number, values in enumerate(rows, start=2):
            if not any(value not in (None, "") for value in values):
                continue
            priority = _positive_integer(
                values[column["priority"]],
                field="priority",
                row_number=row_number,
            )
            vessel_id = _positive_integer(
                values[column["vessel_id"]],
                field="vessel_id",
                row_number=row_number,
            )
            if vessel_id in seen_ids:
                raise ValueError(
                    f"Row {row_number}: duplicate vessel_id {vessel_id}"
                )
            seen_ids.add(vessel_id)
            name = str(values[column["vessel_name"]] or "").strip()
            if not name:
                raise ValueError(f"Row {row_number}: vessel_name is blank")
            link_id = (
                _link_vessel_id(values[column["link"]])
                if "link" in column
                else None
            )
            if link_id is not None and link_id != vessel_id:
                raise ValueError(
                    f"Row {row_number}: link ID {link_id} does not match "
                    f"vessel_id {vessel_id}"
                )
            candidates.append(
                {
                    "input_row": row_number,
                    "priority": priority,
                    "vessel_id": vessel_id,
                    "vessel_name": name,
                    "case": (
                        str(values[column["case"]] or "").strip()
                        if "case" in column
                        else ""
                    ),
                    "edit_url": vessel_edit_url(vessel_id),
                }
            )
        return candidates
    finally:
        workbook.close()


def _selected_candidates(
    candidates: list[dict[str, Any]],
    *,
    priorities: set[int],
    vessel_ids: set[int],
    limit: int | None,
) -> list[dict[str, Any]]:
    selected = [
        candidate
        for candidate in candidates
        if not priorities or candidate["priority"] in priorities
        if not vessel_ids or candidate["vessel_id"] in vessel_ids
    ]
    return selected[:limit] if limit is not None else selected


def _audit_relationship(item: dict[str, Any]) -> dict[str, Any]:
    return {
        "relationship_id": item["relationship_id"],
        "entity_type_id": item["entity_type_id"],
        "owner_name": item["owner_name"],
        "from": item["from"],
        "to": item["to"],
        "status": item["status"],
        "start_date": {
            "year": item["start_year"],
            "month": item["start_month"],
            "day": item["start_day"],
        },
    }


def main() -> int:
    args = build_parser().parse_args()
    if args.limit is not None and args.limit < 1:
        print("--limit must be at least 1", file=sys.stderr)
        return 2
    if any(value < 1 for value in (args.priorities or [])):
        print("--priority values must be positive integers", file=sys.stderr)
        return 2
    if any(value < 1 for value in (args.vessel_ids or [])):
        print("--vessel-id values must be positive integers", file=sys.stderr)
        return 2

    try:
        candidates = load_vessel_candidates(args.input)
    except Exception as exc:
        print(f"Could not load candidate workbook: {exc}", file=sys.stderr)
        return 1
    selected = _selected_candidates(
        candidates,
        priorities=set(args.priorities or []),
        vessel_ids=set(args.vessel_ids or []),
        limit=args.limit,
    )
    if not selected:
        print("No vessels matched the selection.")
        return 0

    stamp = _stamp()
    args.audit_dir.mkdir(parents=True, exist_ok=True)
    audit_path = args.audit_dir / (
        f"vessel-owner-order-{'apply' if args.apply else 'dry-run'}-{stamp}.json"
    )
    audit: dict[str, Any] = {
        "schema_version": 1,
        "dataset": "vessel_owner_order_audit",
        "started_at": utc_now(),
        "mode": "apply" if args.apply else "dry-run",
        "input": str(args.input),
        "filters": {
            "priorities": sorted(set(args.priorities or [])),
            "vessel_ids": sorted(set(args.vessel_ids or [])),
            "limit": args.limit,
        },
        "records": [],
    }
    failures = 0

    try:
        with authenticated_context(headless=args.headless) as context:
            for index, candidate in enumerate(selected, start=1):
                vessel_id = candidate["vessel_id"]
                name = candidate["vessel_name"]
                print(f"[{index}/{len(selected)}] Checking {vessel_id} {name}...")
                record = {**candidate, "status": "pending"}
                audit["records"].append(record)
                try:
                    if candidate["case"].casefold() == "deleted":
                        record["status"] = "skipped_deleted"
                        print("  Skipped because the workbook marks it DELETED.")
                        continue
                    before = fetch_vessel_owner_order(context.session, vessel_id)
                    plan = build_vessel_owner_order_plan(before)
                    record["before"] = [
                        _audit_relationship(item) for item in before
                    ]
                    record["target_relationship_ids"] = plan[
                        "target_relationship_ids"
                    ]
                    record["moves"] = plan["moves"]
                    if not plan["has_changes"]:
                        record["status"] = "no_changes"
                        continue
                    if not args.apply:
                        record["status"] = "planned"
                        print(f"  Planned {len(plan['moves'])} row move(s).")
                        continue

                    submit_vessel_owner_order(
                        context.session,
                        vessel_id=vessel_id,
                        relationship_ids=plan["target_relationship_ids"],
                    )
                    after = fetch_vessel_owner_order(context.session, vessel_id)
                    after_ids = [item["relationship_id"] for item in after]
                    record["after"] = [
                        _audit_relationship(item) for item in after
                    ]
                    if after_ids != plan["target_relationship_ids"]:
                        raise RuntimeError(
                            "Saved owner order did not match the planned order"
                        )
                    record["status"] = "applied_and_verified"
                    print(f"  Applied and verified {len(plan['moves'])} move(s).")
                except Exception as exc:
                    failures += 1
                    record["status"] = "error"
                    record["error"] = str(exc)
                    print(f"  Failed: {exc}", file=sys.stderr)
                finally:
                    audit["updated_at"] = utc_now()
                    atomic_write_json(audit_path, audit)
    except Exception as exc:
        failures += 1
        audit["fatal_error"] = str(exc)
        print(f"Owner-order run stopped: {exc}", file=sys.stderr)

    audit["finished_at"] = utc_now()
    audit["failure_count"] = failures
    audit["processed_count"] = len(audit["records"])
    audit["status_counts"] = dict(
        sorted(Counter(record["status"] for record in audit["records"]).items())
    )
    atomic_write_json(audit_path, audit)
    print(f"Audit report: {audit_path}")
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
